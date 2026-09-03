#!/usr/bin/env python3
"""
import_drugs.py - Automated Drug Catalog Importer with CLI Arguments & Config Support (Issue #22)
"""
import sys
import os
import re
import argparse
import openpyxl
import mysql.connector

def main():
    parser = argparse.ArgumentParser(description="Import pure medical drug price lists into CureRxPro database.")
    parser.add_argument("--excel", default=os.path.join(os.path.dirname(__file__), "Price lists.xlsx"), help="Path to Excel price list file")
    parser.add_argument("--host", default=os.getenv("DB_HOST", "127.0.0.1"), help="MySQL Host")
    parser.add_argument("--port", type=int, default=int(os.getenv("DB_PORT", "3307")), help="MySQL Port")
    parser.add_argument("--user", default=os.getenv("DB_USER", "root"), help="MySQL User")
    parser.add_argument("--password", default=os.getenv("DB_PASS", ""), help="MySQL Password")
    parser.add_argument("--dbname", default=os.getenv("DB_NAME", "prescription_db"), help="MySQL Database Name")
    args = parser.parse_args()

    if not os.path.exists(args.excel):
        print(f"Error: Excel file '{args.excel}' not found.")
        sys.exit(1)

    try:
        conn = mysql.connector.connect(
            host=args.host,
            port=args.port,
            user=args.user,
            password=args.password,
            database=args.dbname
        )
        cursor = conn.cursor()
    except Exception as e:
        print(f"Database connection error: {e}")
        sys.exit(1)

    # Exclusion keywords
    exclude_keywords = [
        'pampers', 'diaper', 'wipes', 'nappy', 'cerelac', 'nan 1', 'nan 2', 'nan 3',
        'lactogen', 'anchor', 'raththi', 'nespray', 'sustagen', 'pediasure',
        'facewash', 'face wash', 'body wash', 'toothbrush', 'brush',
        'whitening', 'fairness', 'cleanser', 'toner', 'night cream', 'day cream',
        'sunscreen', 'sunblock', 'lipstick', 'nail polish', 'hair color', 'hair dye', 'conditioner',
        'perfume', 'deodorant', 'body spray', 'cologne', 'cotton buds', 'durex', 'condom',
        'sanitary', 'whisper', 'stayfree', 'kotex', 'serviette',
        'biscuit', 'chocolate', 'snack', 'water- 5l', 'water 5l', 'water- 1l', 'water- 500ml',
        'candy', 'toffee', 'velvet soap', 'lifebuoy', 'lux soap', 'dove soap', 'pears soap',
        'baby cheramy', 'johnson baby', 'cologne', 'baby powder', 'feeding bottle', 'teat', 'soother',
        'hair oil', 'coconut oil', 'cook oil', 'uno candy', 'wafer', 'touch cologne', 'top dog',
        'sandalwood soap', 'sakura body', 'saffron body', 'avocado body', 'straight conditioner'
    ]

    try:
        print(f"Loading Excel file: {args.excel}...")
        wb = openpyxl.load_workbook(args.excel, data_only=True)
        sheet = wb.active

        insert_query = """
        INSERT INTO `medicines` 
        (`name`, `generic_name`, `type`, `default_strength`, `default_dosage`, `default_instruction`, `unit_price`, `stock_qty`) 
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
        """

        imported_count = 0
        skipped_count = 0
        batch_data = []

        for r in range(2, sheet.max_row + 1):
            name = sheet.cell(r, 2).value
            unit_price = sheet.cell(r, 5).value

            if not name:
                continue

            name_str = str(name).strip()
            lower = name_str.lower()

            if any(kw in lower for kw in exclude_keywords):
                skipped_count += 1
                continue

            try:
                price_val = float(unit_price) if unit_price is not None else 10.00
            except (ValueError, TypeError):
                price_val = 10.00

            strength_match = re.search(r'(\d+(?:\.\d+)?\s*(?:mg|g|ml|mcg|iu|%|IU|MG|ML|G|MCG))', name_str)
            strength = strength_match.group(1).strip() if strength_match else ''

            med_type = 'Tablet'
            default_dosage = '1 Tab tds'
            default_instruction = 'After Meals / කෑමට පසු'

            if any(s in lower for s in ['syrup', 'syr', 'suspension', 'susp', 'liquid', 'solution', 'elixir', 'linctus']):
                med_type = 'Syrup'
                default_dosage = '5ml tds'
            elif any(s in lower for s in ['cream', 'crm', 'ointment', 'gel', 'lotion', 'paste']):
                med_type = 'Cream/Ointment'
                default_dosage = 'Apply thin layer bd'
                default_instruction = 'External Use / බාහිර ආලේපයට'
            elif any(s in lower for s in ['drop', 'eye drop', 'ear drop', 'nasal drop', 'drops']):
                med_type = 'Drops'
                default_dosage = '1-2 Drops tds'
                default_instruction = 'Instill in affected area / යොදන්න'
            elif any(s in lower for s in ['inhaler', 'evohaler', 'rotacap', 'respule', 'spray', 'nasal spray']):
                med_type = 'Inhaler'
                default_dosage = '2 Puffs bd'
                default_instruction = 'Inhale as directed / උරන්න'
            elif any(s in lower for s in ['inj', 'injection', 'vial', 'ampoule']):
                med_type = 'Injection'
                default_dosage = 'As directed'
                default_instruction = 'Administer as directed / එන්නත් කරන්න'
            elif any(s in lower for s in ['cap', 'capsule']):
                med_type = 'Capsule'
                default_dosage = '1 Cap tds'

            batch_data.append((
                name_str,
                name_str,
                med_type,
                strength,
                default_dosage,
                default_instruction,
                round(price_val, 2),
                500
            ))
            imported_count += 1

            if len(batch_data) >= 500:
                cursor.executemany(insert_query, batch_data)
                conn.commit()
                batch_data = []

        if batch_data:
            cursor.executemany(insert_query, batch_data)
            conn.commit()

        cursor.close()
        conn.close()

        print("\n==========================================")
        print(f"IMPORT COMPLETE!")
        print(f"Total Pure Medical Drugs Imported: {imported_count}")
        print(f"Total Non-Medicine / FMCG Items Skipped: {skipped_count}")
        print("==========================================")
    except Exception as e:
        print(f"Error during import execution: {e}")
        sys.exit(1)

if __name__ == "__main__":
    main()
