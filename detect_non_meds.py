import openpyxl
import re

wb = openpyxl.load_workbook('Price lists.xlsx', data_only=True)
sheet = wb['Page1']

# Keywords to detect non-medicine / grocery / cosmetic / baby / FMCG items
non_med_keywords = [
    'pampers', 'diaper', 'baby', 'wipes', 'nappy', 'cerelac', 'nan 1', 'nan 2', 'nan 3',
    'lactogen', 'anchor', 'raththi', 'nespray', 'sustagen', 'ensure', 'pediasure',
    'shampoo', 'soap', 'facewash', 'face wash', 'body wash', 'toothpaste', 'brush',
    'lotion', 'whitening', 'fairness', 'cleanser', 'toner', 'serum', 'night cream', 'day cream',
    'sunscreen', 'sunblock', 'lipstick', 'nail', 'hair color', 'hair dye', 'conditioner',
    'perfume', 'deodorant', 'body spray', 'cologne', 'cotton buds', 'durex', 'condom',
    'sanitary', 'pad', 'whisper', 'stayfree', 'kotex', 'tissue', 'serviette',
    'biscuit', 'chocolate', 'snack', 'drink', 'water', 'glucose', 'candy', 'toffee',
    'oil', 'coconut', 'dettol soap', 'lifebuoy', 'lux', 'dove', 'pears', 'baby cheramy',
    'johnson', 'cologne', 'powder', 'feeding bottle', 'teat', 'soother'
]

non_med_items = []
med_items = []

for r in range(2, sheet.max_row + 1):
    code = sheet.cell(r, 1).value
    name = sheet.cell(r, 2).value
    cost = sheet.cell(r, 3).value
    unit_price = sheet.cell(r, 5).value

    if not name:
        continue

    name_str = str(name).strip()
    lower = name_str.lower()

    matched_keyword = None
    for kw in non_med_keywords:
        if kw in lower:
            matched_keyword = kw
            break

    if matched_keyword:
        non_med_items.append({
            'code': code,
            'name': name_str,
            'price': unit_price,
            'keyword': matched_keyword
        })
    else:
        med_items.append({
            'code': code,
            'name': name_str,
            'price': unit_price
        })

print(f"Total Rows: {sheet.max_row - 1}")
print(f"Detected Non-Medicine / FMCG / Grocery / Cosmetic / Baby items: {len(non_med_items)}")
print(f"Detected Pure Medical / Pharmaceutical items: {len(med_items)}")

print("\nSample Non-Medicine / Cosmetic / Baby / Grocery items found:")
for item in non_med_items[:30]:
    print(f"  [{item['keyword'].upper()}] {item['name']} (Price: LKR {item['price']})")
