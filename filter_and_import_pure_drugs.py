import openpyxl
import re
import mysql.connector

# 1. Connect to MySQL
conn = mysql.connector.connect(
    host='127.0.0.1',
    user='root',
    password='',
    database='prescription_db'
)
cursor = conn.cursor()

cursor.execute("SET FOREIGN_KEY_CHECKS = 0;")
cursor.execute("TRUNCATE TABLE `medicines`;")
cursor.execute("SET FOREIGN_KEY_CHECKS = 1;")
conn.commit()

# 2. Strict Exact-Word and Substring Exclusions for Non-Medicine Items
word_boundary_excludes = [
    r'\btea\b', r'\bbopf\b', r'\bcoffee\b', r'\brice\b', r'\bdhal\b', r'\bsugar\b',
    r'\bsalt\b', r'\bflour\b', r'\bnoodle\b', r'\bnoodles\b', r'\bpasta\b', r'\bmaggi\b',
    r'\bcake\b', r'\bcakes\b', r'\bbiscuit\b', r'\bbiscuits\b', r'\bcookie\b', r'\bcookies\b',
    r'\bcracker\b', r'\bcrackers\b', r'\bwafer\b', r'\bwafers\b', r'\bchocolate\b', r'\bchocolates\b',
    r'\bcandy\b', r'\bcandies\b', r'\btoffee\b', r'\btoffees\b', r'\btofee\b', r'\btofe\b',
    r'\bsweet\b', r'\bsweets\b', r'\bgum\b', r'\bpopcorn\b', r'\bchips\b', r'\bcrisps\b',
    r'\bcashew\b', r'\bpeanut\b', r'\bbutter\b', r'\bcheese\b', r'\bcurd\b', r'\byoghurt\b',
    r'\byogurt\b', r'\bice cream\b', r'\bice pop\b', r'\bjuice\b', r'\bjuices\b', r'\bdrink\b',
    r'\bdrinks\b', r'\bsoda\b', r'\bcoke\b', r'\bpepsi\b', r'\bsprite\b', r'\bfanta\b',
    r'\bsoup\b', r'\bmarmite\b', r'\bseed\b', r'\bseeds\b', r'\bchia\b',
    
    # Diapers / Baby
    r'\bdiaper\b', r'\bdiapers\b', r'\bpampers\b', r'\bnappy\b', r'\bnappies\b', r'\bwipes\b',
    r'\bcerelac\b', r'\blactogen\b', r'\banchor\b', r'\braththi\b', r'\bnespray\b', r'\blakspray\b',
    r'\bsustagen\b', r'\bensure\b', r'\bpediasure\b', r'\bglucerna\b', r'\bmilo\b', r'\bhorlicks\b',
    r'\bsoother\b', r'\bpacifier\b', r'\bfeeding bottle\b', r'\bfeeding cup\b', r'\bteat\b', r'\bteats\b',
    
    # Cosmetics & Toiletries
    r'\bshampoo\b', r'\bconditioner\b', r'\bfacewash\b', r'\bface wash\b', r'\bbody wash\b',
    r'\bbody scrub\b', r'\bface scrub\b', r'\bface pack\b', r'\bface mask\b', r'\bwhitening\b',
    r'\bfairness\b', r'\bcleanser\b', r'\btoner\b', r'\bnight cream\b', r'\bday cream\b',
    r'\bsunscreen\b', r'\bsunblock\b', r'\blipstick\b', r'\blip balm\b', r'\bmascara\b',
    r'\beyeliner\b', r'\bnail polish\b', r'\bhair color\b', r'\bhair dye\b', r'\bhair colour\b',
    r'\bhair gel\b', r'\bhair spray\b', r'\bhair oil\b', r'\bperfume\b', r'\bdeodorant\b',
    r'\bbody spray\b', r'\bcologne\b', r'\bbodymist\b', r'\bbody mist\b', r'\bbody lotion\b',
    r'\bwhisper\b', r'\bstayfree\b', r'\bkotex\b', r'\bfems\b', r'\btampax\b', r'\bpantyliner\b',
    r'\bserviette\b', r'\bserviettes\b', r'\btissue\b', r'\btissues\b', r'\btoilet paper\b',
    r'\bmosquito coil\b', r'\bmosquito mat\b', r'\bmortein\b', r'\bbaygon\b',
    r'\bdurex\b', r'\bcondom\b', r'\bcondoms\b', r'\bmoods\b', r'\bskore\b',
    r'\btoothbrush\b', r'\btoothpaste\b', r'\bsignal\b', r'\bcolgate\b', r'\bclogard\b',
    r'\bharpic\b', r'\blysol\b', r'\bvim\b', r'\bsunlight\b', r'\brinso\b', r'\bsurf excel\b',
    r'\bbattery\b', r'\bbatteries\b', r'\bbulb\b', r'\bbulbs\b',
    r'\btick-out\b', r'\bdog shampoo\b', r'\bcat food\b', r'\bdog food\b',
    
    # Surgical consumables & disposables (to keep pure medical Rx drugs only)
    r'\bsurgical mask\b', r'\bface mask\b', r'\burine bag\b', r'\bneedle\b', r'\bsyringe\b',
    r'\bcannula\b', r'\bcatheter\b', r'\bexamination glove\b', r'\bgloves\b', r'\bcotton buds\b',
    r'\bcotton wool\b', r'\bgauze\b', r'\bplaster\b', r'\bbandage\b', r'\blancet\b', r'\btest strip\b'
]

fmcg_brands = [
    'baby cheramy', 'pears baby', 'panda baby', 'khomba baby', "johnson baby", "johnson's baby",
    'tiki baby', 'wetex', 'velona', 'huggies', 'drypers', 'mamypoko',
    'fair & lovely', 'glow & lovely', 'garnier', "l'oreal", '4ever', 'janet', "nature's secrets",
    'dreamron', 'bellose', 'kumarika', 'keerthi', 'chandanalepa', 'vivya', 'jovees',
    'velvet', 'lux soap', 'lifebuoy', 'dove soap', 'pears soap',
    'dettol soap', 'khomba soap', 'rani soap', 'mysore sandal', 'medimix soap',
    'water- 5l', 'water- 1l', 'water- 500ml', 'water- 1.5l', 'water 5l',
    'swiss roll', 'layer cake', 'sponge cake', 'uno candy', 'waralasha', 'touch cologne',
    'uoodoo', 'uoodo', 'tiara swiss', 'tiara sponge', 'thalawakele', 'vanilla layer', 'vanila cup',
    'wet look hair', 'uswaththa', 'watawala', 'alpenliebe', 'mentos', 'eclairs', 'halls'
]

compiled_regexes = [re.compile(p, re.IGNORECASE) for p in word_boundary_excludes]

print("Reading and strictly filtering Price lists.xlsx...")
wb = openpyxl.load_workbook('Price lists.xlsx', data_only=True)
sheet = wb['Page1']

insert_query = """
INSERT INTO `medicines` 
(`name`, `generic_name`, `type`, `default_strength`, `default_dosage`, `default_instruction`, `unit_price`, `stock_qty`) 
VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
"""

imported_count = 0
skipped_count = 0
batch_data = []

for r in range(2, sheet.max_row + 1):
    code = sheet.cell(r, 1).value
    name = sheet.cell(r, 2).value
    unit_price = sheet.cell(r, 5).value

    if not name:
        continue

    name_str = str(name).strip()
    lower = name_str.lower()

    # 1. Check FMCG brand substrings
    is_fmcg = False
    for b in fmcg_brands:
        if b in lower:
            is_fmcg = True
            break

    # 2. Check word boundary regexes
    if not is_fmcg:
        for rgx in compiled_regexes:
            if rgx.search(lower):
                is_fmcg = True
                break

    if is_fmcg:
        skipped_count += 1
        continue

    # Parse Unit Price
    try:
        price_val = float(unit_price) if unit_price is not None else 10.00
    except (ValueError, TypeError):
        price_val = 10.00

    # Extract strength
    strength_match = re.search(r'(\d+(?:\.\d+)?\s*(?:mg|g|ml|mcg|iu|%|IU|MG|ML|G|MCG))', name_str)
    strength = strength_match.group(1).strip() if strength_match else ''

    # Determine Type, Default Dosage & Instruction
    med_type = 'Tablet'
    default_dosage = '1 Tab tds'
    default_instruction = 'After Meals / කෑමට පසු'

    if any(s in lower for s in ['syrup', 'syr', 'suspension', 'susp', 'liquid', 'solution', 'elixir', 'linctus']):
        med_type = 'Syrup'
        default_dosage = '5ml tds'
        default_instruction = 'After Meals / කෑමට පසු'
    elif any(s in lower for s in ['cream', 'crm', 'ointment', 'gel', 'lotion', 'paste']):
        med_type = 'Cream/Ointment'
        default_dosage = 'Apply thin layer bd'
        default_instruction = 'External Use / බාහිර ආලේපයට'
    elif any(s in lower for s in ['drop', 'eye drop', 'ear drop', 'nasal drop', 'drops']):
        med_type = 'Drops'
        default_dosage = '1-2 Drops tds'
        default_instruction = 'Instill in affected area / යොදන්න'
    elif any(s in lower for s in ['inhaler', 'evohaler', 'rotacap', 'respule', 'spray', 'nasal spray']):
        med_type = 'Inhaler'
        default_dosage = '2 Puffs bd'
        default_instruction = 'Inhale as directed / උරන්න'
    elif any(s in lower for s in ['inj', 'injection', 'vial', 'ampoule']):
        med_type = 'Injection'
        default_dosage = 'As directed'
        default_instruction = 'Administer as directed / එන්නත් කරන්න'
    elif any(s in lower for s in ['cap', 'capsule']):
        med_type = 'Capsule'
        default_dosage = '1 Cap tds'
        default_instruction = 'After Meals / කෑමට පසු'
    else:
        med_type = 'Tablet'
        default_dosage = '1 Tab tds'
        default_instruction = 'After Meals / කෑමට පසු'

    generic_name = name_str

    batch_data.append((
        name_str,
        generic_name,
        med_type,
        strength,
        default_dosage,
        default_instruction,
        round(price_val, 2),
        500
    ))
    imported_count += 1

    if len(batch_data) >= 500:
        cursor.executemany(insert_query, batch_data)
        conn.commit()
        batch_data = []

if batch_data:
    cursor.executemany(insert_query, batch_data)
    conn.commit()

cursor.close()
conn.close()

print(f"\n==========================================")
print(f"STRICT MEDICAL IMPORT COMPLETED!")
print(f"Total Pure Medical Drugs In Database: {imported_count}")
print(f"Total Grocery / Cosmetic / Baby / Surgical Excluded: {skipped_count}")
print(f"==========================================")
