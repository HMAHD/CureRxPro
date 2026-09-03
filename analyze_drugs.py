import openpyxl
import re

wb = openpyxl.load_workbook('Price lists.xlsx', data_only=True)
sheet = wb['Page1']

total = 0
valid_drugs = 0
types_count = {}

for r in range(2, sheet.max_row + 1):
    code = sheet.cell(r, 1).value
    name = sheet.cell(r, 2).value
    cost = sheet.cell(r, 3).value
    unit_price = sheet.cell(r, 5).value

    if not name:
        continue

    total += 1
    name_str = str(name).strip()

    # Extract strength (e.g. 500mg, 50mg, 10ml, 5%, 200iu)
    strength_match = re.search(r'(\d+(?:\.\d+)?\s*(?:mg|g|ml|mcg|iu|%|IU|MG|ML|G|MCG))', name_str)
    strength = strength_match.group(1) if strength_match else ''

    # Determine type
    lower_name = name_str.lower()
    med_type = 'Tablet'
    if 'syrup' in lower_name or 'syr' in lower_name or 'suspension' in lower_name or 'susp' in lower_name or 'liquid' in lower_name:
        med_type = 'Syrup'
    elif 'cream' in lower_name or 'crm' in lower_name or 'ointment' in lower_name or 'gel' in lower_name:
        med_type = 'Cream/Ointment'
    elif 'drop' in lower_name:
        med_type = 'Drops'
    elif 'inhaler' in lower_name or 'evohaler' in lower_name or 'rotacap' in lower_name:
        med_type = 'Inhaler'
    elif 'inj' in lower_name or 'injection' in lower_name:
        med_type = 'Injection'
    elif 'cap' in lower_name or 'capsule' in lower_name:
        med_type = 'Capsule'
    elif 'tablet' in lower_name or 'tab' in lower_name:
        med_type = 'Tablet'

    types_count[med_type] = types_count.get(med_type, 0) + 1

print(f"Total Valid Drug Records: {total}")
print("Detected Types Breakdown:")
for k, v in sorted(types_count.items(), key=lambda x: x[1], reverse=True):
    print(f"  - {k}: {v}")

print("\nFirst 10 sample parsed items:")
for r in range(2, 12):
    name = str(sheet.cell(r, 2).value).strip()
    uprice = sheet.cell(r, 5).value
    print(f"  * Product: {name} -> Unit Price: LKR {uprice}")
